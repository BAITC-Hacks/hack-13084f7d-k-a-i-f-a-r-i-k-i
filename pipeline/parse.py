"""Шаг 1. Парсинг Word / PDF / Excel в список Chunk.   Владелец: роль 2.

Базовая версия уже работает на data/test_set. TODO:
- функции таблицей внутри Word (doc.tables)
- многоуровневая нумерация 3.1.1, пункты вида «1)», «а)»
- сканы PDF (OCR) — только если организаторы дадут сканы
"""
import re
from pathlib import Path
from .schemas import Chunk

CLAUSE_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(.+)$", re.S)
SECTION_RE = re.compile(r"^\s*(\d+)\.\s+([А-ЯЁA-Z][^.]{2,80})$")   # «3. Функции»


def _lines_to_chunks(lines, doc_name, side, pages=None) -> list[Chunk]:
    """lines: список строк-абзацев; pages: параллельный список номеров страниц."""
    chunks, section, cur = [], "", None
    pages = pages or [None] * len(lines)
    for line, page in zip(lines, pages):
        line = line.strip()
        if not line:
            continue
        sec = SECTION_RE.match(line)
        if sec and "." not in sec.group(1):
            section = line
            cur = None
            continue
        m = CLAUSE_RE.match(line)
        if m and "." in m.group(1):          # пункт вида 3.2
            cur = {"chunk_id": f"{side}/{doc_name}#{m.group(1)}", "doc_name": doc_name, "side": side,
                   "page": page, "clause": m.group(1), "section": section, "text": m.group(2).strip()}
            chunks.append(cur)
        elif cur is not None and not m:      # продолжение пункта, перенесённое на новую строку (PDF)
            cur["text"] += " " + line
        else:                                # шапка / свободный текст / пункт приказа «1. ...»
            clause = m.group(1) if m else f"p{len(chunks)}"
            text = m.group(2).strip() if m else line
            cur = {"chunk_id": f"{side}/{doc_name}#{clause}", "doc_name": doc_name, "side": side,
                   "page": page, "clause": clause, "section": section or "header", "text": text}
            chunks.append(cur)
    # уникальность chunk_id
    seen = {}
    for c in chunks:
        k = c["chunk_id"]
        if k in seen:
            seen[k] += 1
            c["chunk_id"] = f"{k}~{seen[k]}"
        else:
            seen[k] = 0
    return chunks


def parse_docx(path: Path, side) -> list[Chunk]:
    from docx import Document
    d = Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:                        # таблицы — построчно
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return _lines_to_chunks(lines, path.name, side)


def parse_pdf(path: Path, side) -> list[Chunk]:
    import pymupdf as fitz
    lines, pages = [], []
    with fitz.open(str(path)) as doc:
        for pno, page in enumerate(doc, 1):
            for ln in page.get_text("text").splitlines():
                lines.append(ln); pages.append(pno)
    return _lines_to_chunks(lines, path.name, side, pages)


def parse_xlsx(path: Path, side) -> list[Chunk]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        header = None
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if not vals:
                continue
            if header is None and len(vals) >= 3:
                header = [str(v).strip() if v is not None else "" for v in row]
                continue
            if header:
                pairs = [f"{h}: {v}" for h, v in zip(header, row) if h and v is not None]
                text = "; ".join(pairs)
            else:
                text = " ".join(vals)
            out.append({"chunk_id": f"{side}/{path.name}#{ws.title}!row:{r_idx}", "doc_name": path.name,
                        "side": side, "page": None, "clause": f"row:{r_idx}", "section": ws.title, "text": text})
    return out


PARSERS = {".docx": parse_docx, ".pdf": parse_pdf, ".xlsx": parse_xlsx}


def parse_file(path, side) -> list[Chunk]:
    path = Path(path)
    fn = PARSERS.get(path.suffix.lower())
    if not fn:
        raise ValueError(f"Неподдерживаемый формат: {path.name}")
    return fn(path, side)


def parse_folder(folder, side) -> list[Chunk]:
    chunks = []
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() in PARSERS:
            chunks += parse_file(p, side)
    return chunks
